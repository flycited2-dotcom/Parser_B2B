"""Build an approval-gated workbook for the existing email automation.

The canonical master remains complete.  This module creates a smaller,
auditable delivery candidate file with the exact Russian headers consumed by
``Email_horeca_send``.  It never sends mail and never marks data as approved.
"""
from __future__ import annotations

import csv
import os
import re
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from utils.csv_safety import neutralize_csv_formula
from utils.quality import VK_QUARANTINE_FLAGS


OUTREACH_HEADERS = [
    "Email", "Название", "Тип клиента", "Город", "Телефон", "Сайт",
    "Соцсеть", "Адрес", "Источник", "ID объекта", "Доверие",
    "Флаги качества", "Все email", "Все телефоны",
]

TARGET_TYPES = {
    "ресторан", "кафе", "фастфуд", "бар", "паб", "клуб", "кофейня",
    "столовая", "фудкорт", "пиццерия", "кондитерская",
}

HARD_REJECT_FLAGS = {"outside_crimea", "vk_negative_terms", *VK_QUARANTINE_FLAGS}
MANUAL_SEGMENT_FLAGS = {"vk_manual_business_segment"}
EMAIL_RE = re.compile(r"[a-z0-9._%+\-]+@[a-z0-9.\-]+\.[a-z]{2,}", re.I)
EMAIL_BLACKLIST = re.compile(
    r"(?:^|@)(?:example(?:\.|@)|test(?:\.|@))|"
    r"no-?reply|mailer-daemon|postmaster@|webmaster@|abuse@|"
    r"@.*\.gov\.ru$|@stacks\.vk-portal\.net$",
    re.I,
)
PREFERRED_PREFIXES = (
    "reservation", "reservations", "booking", "reception", "info",
    "order", "zakaz", "sales", "manager", "office", "contact",
)


def _confidence(row: dict) -> float:
    try:
        value = float(str(row.get("confidence") or ""))
    except ValueError:
        return 0.0
    return max(0.0, min(1.0, value))


def _flags(row: dict) -> set[str]:
    return {
        flag.strip().casefold()
        for flag in str(row.get("quality_flags") or "").split("|")
        if flag.strip()
    }


def _website_domain(value: str) -> str:
    raw = str(value or "").strip()
    if not raw:
        return ""
    try:
        host = (urlparse(raw if "://" in raw else "https://" + raw).hostname or "")
        return host.casefold().removeprefix("www.").rstrip(".")
    except ValueError:
        return ""


def _email_candidates(row: dict) -> list[str]:
    website_domain = _website_domain(row.get("website", ""))
    candidates: dict[str, tuple[int, str]] = {}
    material = " | ".join(
        str(row.get(field) or "") for field in ("email", "all_emails")
    )
    for match in EMAIL_RE.findall(material):
        email = match.strip().casefold()
        if EMAIL_BLACKLIST.search(email):
            continue
        local, _separator, domain = email.partition("@")
        score = 0
        if website_domain and (
            domain == website_domain
            or domain.endswith("." + website_domain)
            or website_domain.endswith("." + domain)
        ):
            score += 100
        for index, prefix in enumerate(PREFERRED_PREFIXES):
            if local == prefix or local.startswith(prefix + "."):
                score += 50 - index
                break
        current = candidates.get(email)
        if current is None or score > current[0]:
            candidates[email] = (score, email)
    return [
        email
        for _score, email in sorted(
            candidates.values(), key=lambda item: (-item[0], item[1])
        )
    ]


def _review_reason(row: dict, min_confidence: float) -> str:
    if not str(row.get("name") or "").strip():
        return "missing_name"
    if str(row.get("client_type") or "").strip().casefold() not in TARGET_TYPES:
        return "non_target_type"
    emails = _email_candidates(row)
    if not emails:
        return "missing_or_invalid_email"
    flags = _flags(row)
    if flags & HARD_REJECT_FLAGS:
        return "hard_quality_flag"
    if "vk_weak_contact_donor" in flags:
        return "weak_vk_contact_donor"
    if flags & MANUAL_SEGMENT_FLAGS:
        return "manual_business_segment"
    confidence = _confidence(row)
    if confidence < min_confidence:
        return "low_confidence"
    if "manual_review" in flags and confidence < max(min_confidence, 0.85):
        return "manual_review"
    return ""


def _outreach_row(row: dict, email: str) -> dict:
    return {
        "Email": email,
        "Название": row.get("name", ""),
        "Тип клиента": row.get("client_type", ""),
        "Город": row.get("city", ""),
        "Телефон": row.get("phone", ""),
        "Сайт": row.get("website", ""),
        "Соцсеть": row.get("social", ""),
        "Адрес": row.get("address", ""),
        "Источник": row.get("sources") or row.get("source", ""),
        "ID объекта": row.get("entity_id", ""),
        "Доверие": f"{_confidence(row):.2f}",
        "Флаги качества": row.get("quality_flags", ""),
        "Все email": row.get("all_emails") or row.get("email", ""),
        "Все телефоны": row.get("all_phones") or row.get("phone", ""),
    }


def _safe_row(row: dict, headers: list[str]) -> dict:
    return {
        header: neutralize_csv_formula(str(row.get(header) or ""))
        for header in headers
    }


def _atomic_csv(path: Path, headers: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    with tmp.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=headers,
            delimiter=";",
            quoting=csv.QUOTE_ALL,
            extrasaction="ignore",
        )
        writer.writeheader()
        writer.writerows(_safe_row(row, headers) for row in rows)
    os.replace(tmp, path)


def _atomic_xlsx(
    path: Path,
    rows: list[dict],
    *,
    generated_at: str,
    run_id: str,
    min_confidence: float,
    review_count: int,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp.xlsx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Контакты"
    header_fill = PatternFill("solid", fgColor="1F2937")
    for column, header in enumerate(OUTREACH_HEADERS, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for row_number, row in enumerate(rows, start=2):
        safe = _safe_row(row, OUTREACH_HEADERS)
        for column, header in enumerate(OUTREACH_HEADERS, start=1):
            sheet.cell(row=row_number, column=column, value=safe[header])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:N{max(1, len(rows) + 1)}"
    widths = [30, 36, 18, 16, 22, 34, 30, 40, 18, 30, 12, 30, 44, 44]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    metadata = workbook.create_sheet("Метаданные")
    metadata_rows = [
        ("schema_version", 1),
        ("state", "ready_for_review"),
        ("approved_for_send", False),
        ("run_id", run_id),
        ("generated_at", generated_at),
        ("ready_rows", len(rows)),
        ("review_rows", review_count),
        ("min_confidence", min_confidence),
        ("consumer", "Email_horeca_send"),
    ]
    for row_number, (key, value) in enumerate(metadata_rows, start=1):
        metadata.cell(row=row_number, column=1, value=key).font = Font(bold=True)
        metadata.cell(row=row_number, column=2, value=value)
    metadata.column_dimensions["A"].width = 24
    metadata.column_dimensions["B"].width = 38
    workbook.save(tmp)
    os.replace(tmp, path)


def build_outreach_exports(
    master_csv: str,
    output_dir: str | None = None,
    *,
    run_id: str = "",
    min_confidence: float | None = None,
) -> dict:
    """Create ready/review artifacts and return their paths and counts."""
    source = Path(master_csv)
    target_dir = Path(output_dir) if output_dir else source.parent
    if min_confidence is None:
        try:
            min_confidence = float(os.getenv("OUTREACH_MIN_CONFIDENCE", "0.70"))
        except ValueError:
            min_confidence = 0.70
    min_confidence = max(0.0, min(1.0, min_confidence))

    with source.open(newline="", encoding="utf-8-sig") as handle:
        master_rows = list(csv.DictReader(handle, delimiter=";"))

    candidates: list[tuple[float, int, dict, str]] = []
    review_rows: list[dict] = []
    for row in master_rows:
        reason = _review_reason(row, min_confidence)
        if reason:
            review_rows.append({**row, "review_reason": reason})
            continue
        email = _email_candidates(row)[0]
        completeness = sum(bool(row.get(field)) for field in (
            "phone", "website", "social", "address",
        ))
        candidates.append((_confidence(row), completeness, row, email))

    ready_rows: list[dict] = []
    used_emails: set[str] = set()
    for _confidence_value, _completeness, row, email in sorted(
        candidates,
        key=lambda item: (-item[0], -item[1], item[3]),
    ):
        if email in used_emails:
            review_rows.append({**row, "review_reason": "duplicate_email"})
            continue
        used_emails.add(email)
        ready_rows.append(_outreach_row(row, email))
    ready_rows.sort(key=lambda row: (str(row["Город"]), str(row["Название"]), row["Email"]))

    ready_csv = target_dir / "outreach_ready.csv"
    ready_xlsx = target_dir / "outreach_ready.xlsx"
    review_headers = list(master_rows[0].keys()) + ["review_reason"] if master_rows else ["review_reason"]
    review_csv = target_dir / "outreach_review.csv"
    generated_at = datetime.now().astimezone().isoformat(timespec="seconds")
    _atomic_csv(ready_csv, OUTREACH_HEADERS, ready_rows)
    _atomic_csv(review_csv, review_headers, review_rows)
    _atomic_xlsx(
        ready_xlsx,
        ready_rows,
        generated_at=generated_at,
        run_id=run_id,
        min_confidence=min_confidence,
        review_count=len(review_rows),
    )
    return {
        "ready_csv": str(ready_csv),
        "ready_xlsx": str(ready_xlsx),
        "review_csv": str(review_csv),
        "ready_rows": len(ready_rows),
        "review_rows": len(review_rows),
        "min_confidence": min_confidence,
        "approved_for_send": False,
    }
