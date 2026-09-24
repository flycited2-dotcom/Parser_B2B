import csv
import json

from openpyxl import load_workbook

from utils.outreach_export import OUTREACH_HEADERS, build_outreach_exports
from utils.storage import FIELDS


def _write_master(path, rows):
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def _provenance(*segments):
    return json.dumps({"observations": [], "fields": {
        "client_type": [{"value": key, "source": "OSM", "observation_id": f"o{i}"}
                        for i, key in enumerate(segments)],
    }}, ensure_ascii=False)


ROWS = [
    {"entity_id": "ok", "name": "Окна Юг", "city": "Ялта", "client_type": "stroitelstvo",
     "email": "info@okna-yug.ru", "website": "https://okna-yug.ru", "confidence": "0.95",
     "sources": "OSM|Я.Карты", "provenance": _provenance("stroitelstvo", "mebel")},
    {"entity_id": "nosite", "name": "Салон Лилия", "city": "Симферополь", "client_type": "krasota",
     "email": "lilia@mail.ru", "social": "https://vk.com/lilia", "confidence": "0.90"},
    {"entity_id": "chain", "name": "Магнит Косметик", "city": "Ялта", "client_type": "krasota",
     "email": "info@beauty.ru", "confidence": "0.99", "quality_flags": "excluded_chain"},
    {"entity_id": "chainmail", "name": "Салон связи", "city": "Ялта", "client_type": "torgovlya",
     "email": "shop@mts.ru", "confidence": "0.99"},
    {"entity_id": "competitor", "name": "Веб-студия Код", "city": "Ялта", "client_type": "it_svyaz",
     "email": "hi@kod.ru", "confidence": "0.99"},
    {"entity_id": "nosegment", "name": "ООО Ромашка", "city": "Ялта", "client_type": "прочее",
     "email": "info@romashka.ru", "confidence": "0.99"},
    {"entity_id": "otherbase", "name": "Кейтеринг Юг", "city": "Ялта", "client_type": "event",
     "email": "info@cafe.ru", "confidence": "0.90"},
    {"entity_id": "dup", "name": "Окна Юг филиал", "city": "Ялта", "client_type": "stroitelstvo",
     "email": "info@okna-yug.ru", "confidence": "0.80"},
    {"entity_id": "low", "name": "Сомнительно", "city": "Керчь", "client_type": "avto",
     "email": "a@maybe.ru", "confidence": "0.40"},
    {"entity_id": "noise", "name": "Барахолка", "city": "Ялта", "client_type": "torgovlya",
     "email": "b@bar.ru", "confidence": "0.90", "quality_flags": "vk_noise_primary"},
]
CACHE = {"okna-yug.ru": {"signals": ["site_builder", "no_https"], "checked_at": "2026-09-24T00:00:00+00:00"}}


def _read(path):
    with open(path, encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle, delimiter=";"))


def test_outreach_rows_segments_signals_and_review_reasons(tmp_path):
    master = tmp_path / "master_all.csv"
    _write_master(master, ROWS)

    result = build_outreach_exports(
        str(master), str(tmp_path), run_id="t", min_confidence=0.7,
        signals_cache=CACHE, other_base=({"info@cafe.ru"}, set()),
        include_competitors=False,
    )

    assert result["ready_rows"] == 2
    assert result["review_rows"] == 8
    assert result["approved_for_send"] is False
    assert result["by_segment"] == {"stroitelstvo": 1, "krasota": 1}
    assert result["by_signal"] == {"no_https": 1, "site_builder": 1, "no_website": 1}

    ready = {row["Email"]: row for row in _read(result["ready_csv"])}
    assert list(next(iter(ready.values()))) == OUTREACH_HEADERS
    okna = ready["info@okna-yug.ru"]
    assert okna["Сегмент"] == "Строительство и ремонт"
    assert okna["Все сегменты"] == "Строительство и ремонт; Мебель и интерьер"
    assert okna["Сигналы"] == "no_https; site_builder"
    assert okna["Повод для КП"] == "Модернизация сайта"
    lilia = ready["lilia@mail.ru"]
    assert lilia["Сигналы"] == "no_website"
    assert lilia["Повод для КП"] == "Сайт с нуля"

    reasons = {row["entity_id"]: row["review_reason"] for row in _read(result["review_csv"])}
    assert reasons == {
        "chain": "excluded_chain",
        "chainmail": "missing_or_invalid_email",
        "competitor": "competitor_segment",
        "nosegment": "no_segment",
        "otherbase": "already_in_other_base",
        "dup": "duplicate_email",
        "low": "low_confidence",
        "noise": "hard_quality_flag",
    }

    workbook = load_workbook(result["ready_xlsx"], read_only=True)
    assert workbook.sheetnames == ["Контакты", "Метаданные"]
    assert [cell.value for cell in next(workbook["Контакты"].iter_rows())] == OUTREACH_HEADERS
    metadata = dict(workbook["Метаданные"].iter_rows(values_only=True))
    assert metadata["approved_for_send"] is False
    assert metadata["consumer"] == "manual_review_only"


def test_competitors_can_be_included_explicitly(tmp_path):
    master = tmp_path / "master_all.csv"
    _write_master(master, [ROWS[4]])
    result = build_outreach_exports(
        str(master), str(tmp_path), min_confidence=0.7, include_competitors=True,
    )
    assert result["ready_rows"] == 1
    row = _read(result["ready_csv"])[0]
    assert row["Сигналы"] == "no_website"


def test_exclusions_are_recomputed_so_list_updates_apply_to_stored_rows(tmp_path):
    master = tmp_path / "master_all.csv"
    _write_master(master, [
        {"entity_id": "chain-unflagged", "name": "Магнит Косметик", "city": "Ялта",
         "client_type": "krasota", "email": "info@beauty-shop.ru", "confidence": "0.99"},
        {"entity_id": "gov-in-provenance", "name": "Стоматология", "city": "Ялта",
         "client_type": "medicina", "email": "info@stom.ru", "confidence": "0.99",
         "provenance": json.dumps({"observations": [], "fields": {"name": [
             {"value": "ГАУЗРК Стоматологическая поликлиника №1", "source": "OSM", "observation_id": "o1"},
         ]}}, ensure_ascii=False)},
    ])
    result = build_outreach_exports(str(master), str(tmp_path), min_confidence=0.7)
    assert result["ready_rows"] == 0
    reasons = {row["entity_id"]: row["review_reason"] for row in _read(result["review_csv"])}
    assert reasons == {"chain-unflagged": "excluded_chain", "gov-in-provenance": "excluded_gov"}
