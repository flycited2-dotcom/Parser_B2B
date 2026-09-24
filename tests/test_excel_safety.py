import pytest
from openpyxl import Workbook

from utils.csv_safety import neutralize_csv_formula
from utils.excel_export import _write_sheet


@pytest.mark.parametrize("value", [
    "=1+1",
    "+cmd|' /C calc'!A0",
    "-2+3",
    "@SUM(A1:A2)",
    "  =HYPERLINK(\"http://example.test\")",
    "\t@SUM(A1:A2)",
    "\n=1+1",
])
def test_formula_like_values_are_neutralized(value):
    safe = neutralize_csv_formula(value)
    assert safe.startswith("'")
    assert safe[1:] == value


def test_normal_text_is_unchanged():
    assert neutralize_csv_formula("Кафе Лето") == "Кафе Лето"
    assert neutralize_csv_formula(42) == "42"


def test_xlsx_cells_are_strings_not_formulas():
    wb = Workbook()
    ws = wb.active
    _write_sheet(ws, [{
        "city": "=1+1",
        "name": "+cmd|' /C calc'!A0",
        "phone": "+7 (999) 123-45-67",
    }])
    assert ws["A2"].data_type == "s"
    assert ws["A2"].value == "'=1+1"
    assert ws["B2"].data_type == "s"
    assert ws["B2"].value.startswith("'+cmd")
    assert ws["F2"].data_type == "s"
    assert ws["F2"].value.startswith("'+7")
